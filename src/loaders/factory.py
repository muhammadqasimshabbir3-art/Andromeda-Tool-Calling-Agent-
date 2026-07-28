"""Load PDF, images, text, and Office documents into a common representation."""

from __future__ import annotations

import base64
import hashlib
import mimetypes
from io import BytesIO
from pathlib import Path

from config.settings import get_settings
from loaders.base import LoadedDocument, LoadedPage, LoaderError

SUPPORTED_IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    ".gif",
}
SUPPORTED_PDF_EXTENSIONS = {".pdf"}
SUPPORTED_TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".log", ".text"}
SUPPORTED_DOCX_EXTENSIONS = {".docx"}

SUPPORTED_EXTENSIONS = (
    SUPPORTED_PDF_EXTENSIONS
    | SUPPORTED_IMAGE_EXTENSIONS
    | SUPPORTED_TEXT_EXTENSIONS
    | SUPPORTED_DOCX_EXTENSIONS
)


def decode_document_payload(data_base64: str) -> bytes:
    """Decode and validate a base64 document payload from the UI."""
    settings = get_settings()
    if not data_base64:
        raise LoaderError("No document file was provided.")
    try:
        payload = base64.b64decode(data_base64, validate=True)
    except Exception as exc:  # noqa: BLE001
        raise LoaderError("The uploaded document payload is invalid.") from exc
    if not payload:
        raise LoaderError("The uploaded document file is empty.")
    if len(payload) > settings.max_document_bytes:
        limit_mb = settings.max_document_bytes // (1024 * 1024)
        raise LoaderError(f"The uploaded document is too large (limit {limit_mb} MB).")
    return payload


def _guess_mime(filename: str, payload: bytes) -> str:
    name = (filename or "").lower()
    guessed, _ = mimetypes.guess_type(name)
    if guessed:
        return guessed
    if payload.startswith(b"%PDF"):
        return "application/pdf"
    if payload.startswith(b"\x89PNG"):
        return "image/png"
    if payload[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if payload[:2] == b"BM":
        return "image/bmp"
    if payload[:4] in {b"II*\x00", b"MM\x00*"}:
        return "image/tiff"
    if payload[:4] == b"RIFF" and payload[8:12] == b"WEBP":
        return "image/webp"
    if payload[:6] in {b"GIF87a", b"GIF89a"}:
        return "image/gif"
    if payload[:2] == b"PK" and name.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return "application/octet-stream"


def _load_pdf(payload: bytes, filename: str, fingerprint: str) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise LoaderError("PDF loading requires the `pypdf` package.") from exc

    try:
        reader = PdfReader(BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        raise LoaderError("The uploaded PDF could not be read.") from exc

    if getattr(reader, "is_encrypted", False):
        raise LoaderError("Encrypted or password-protected PDF files are not supported.")

    pages_obj = getattr(reader, "pages", [])
    if not pages_obj:
        raise LoaderError("The uploaded PDF does not contain any pages.")

    pages: list[LoadedPage] = []
    for page_number, page in enumerate(pages_obj, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            page_text = ""
        pages.append(
            LoadedPage(
                page_number=page_number,
                text=page_text,
                metadata={"source": "pypdf"},
            )
        )

    return LoadedDocument(
        filename=filename or "uploaded.pdf",
        mime_type="application/pdf",
        fingerprint=fingerprint,
        pages=pages,
        raw_bytes=payload,
        metadata={"page_count": len(pages)},
    )


def _load_image(payload: bytes, filename: str, fingerprint: str, mime_type: str) -> LoadedDocument:
    try:
        from PIL import Image
    except ImportError as exc:
        raise LoaderError("Image loading requires the `Pillow` package.") from exc

    try:
        image = Image.open(BytesIO(payload))
        image.load()
        width, height = image.size
    except Exception as exc:  # noqa: BLE001
        raise LoaderError("The uploaded image could not be read.") from exc

    page = LoadedPage(
        page_number=1,
        text="",
        image_bytes=payload,
        width=width,
        height=height,
        metadata={"source": "pillow", "mode": getattr(image, "mode", "")},
    )
    return LoadedDocument(
        filename=filename or "uploaded.png",
        mime_type=mime_type,
        fingerprint=fingerprint,
        pages=[page],
        raw_bytes=payload,
        metadata={"page_count": 1},
    )


def _decode_text_bytes(payload: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _pages_from_long_text(text: str, *, source: str, page_chars: int = 3500) -> list[LoadedPage]:
    cleaned = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not cleaned:
        raise LoaderError("The uploaded text file is empty.")
    pages: list[LoadedPage] = []
    for index in range(0, len(cleaned), page_chars):
        chunk = cleaned[index : index + page_chars].strip()
        if not chunk:
            continue
        pages.append(
            LoadedPage(
                page_number=len(pages) + 1,
                text=chunk,
                metadata={"source": source},
            )
        )
    if not pages:
        raise LoaderError("No readable text could be extracted from the file.")
    return pages


def _load_text(payload: bytes, filename: str, fingerprint: str, mime_type: str) -> LoadedDocument:
    text = _decode_text_bytes(payload)
    pages = _pages_from_long_text(text, source="text_file")
    return LoadedDocument(
        filename=filename or "uploaded.txt",
        mime_type=mime_type or "text/plain",
        fingerprint=fingerprint,
        pages=pages,
        raw_bytes=payload,
        metadata={"page_count": len(pages), "format": "text"},
    )


def _load_docx(payload: bytes, filename: str, fingerprint: str) -> LoadedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise LoaderError("DOCX loading requires the `python-docx` package.") from exc

    try:
        document = Document(BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        raise LoaderError("The uploaded Word document (.docx) could not be read.") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        body = (paragraph.text or "").strip()
        if body:
            parts.append(body)
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)

    text = "\n".join(parts).strip()
    if not text:
        raise LoaderError("The Word document has no extractable text.")

    pages = _pages_from_long_text(text, source="python-docx")
    return LoadedDocument(
        filename=filename or "uploaded.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        fingerprint=fingerprint,
        pages=pages,
        raw_bytes=payload,
        metadata={"page_count": len(pages), "format": "docx"},
    )


def load_document(
    data_base64: str,
    filename: str = "uploaded.bin",
    mime_type: str | None = None,
) -> LoadedDocument:
    """Decode base64 input and return a LoadedDocument."""
    payload = decode_document_payload(data_base64)
    fingerprint = hashlib.sha256(payload).hexdigest()
    name = filename or "uploaded.bin"
    ext = Path(name).suffix.lower()
    resolved_mime = mime_type or _guess_mime(name, payload)

    if (
        ext in SUPPORTED_PDF_EXTENSIONS
        or resolved_mime == "application/pdf"
        or payload.startswith(b"%PDF")
    ):
        return _load_pdf(payload, name, fingerprint)

    if ext in SUPPORTED_IMAGE_EXTENSIONS or resolved_mime.startswith("image/"):
        return _load_image(payload, name, fingerprint, resolved_mime)

    if ext in SUPPORTED_TEXT_EXTENSIONS or resolved_mime.startswith("text/"):
        return _load_text(payload, name, fingerprint, resolved_mime)

    if (
        ext in SUPPORTED_DOCX_EXTENSIONS
        or resolved_mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return _load_docx(payload, name, fingerprint)

    supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
    raise LoaderError(
        f"Unsupported file type ({ext or resolved_mime}). "
        f"Supported: {supported}"
    )
