"""Unit tests for document loaders."""

import base64
from io import BytesIO

import pytest

from loaders.base import LoaderError
from loaders.factory import decode_document_payload, load_document


def test_decode_rejects_empty():
    with pytest.raises(LoaderError):
        decode_document_payload("")


def test_load_png_image():
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (32, 32), color=(255, 255, 255)).save(buf, format="PNG")
    payload = base64.b64encode(buf.getvalue()).decode("ascii")
    doc = load_document(payload, filename="scan.png")
    assert doc.mime_type == "image/png"
    assert doc.needs_ocr is True
    assert len(doc.pages) == 1


def test_load_gif_image():
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (24, 24), color=(10, 20, 30)).save(buf, format="GIF")
    payload = base64.b64encode(buf.getvalue()).decode("ascii")
    doc = load_document(payload, filename="scan.gif", mime_type="image/gif")
    assert doc.mime_type == "image/gif"
    assert doc.needs_ocr is True


def test_load_plain_text_arabic():
    text = "هذا نص عربي للاختبار.\n" * 5
    payload = base64.b64encode(text.encode("utf-8")).decode("ascii")
    doc = load_document(payload, filename="notes.txt")
    assert doc.mime_type.startswith("text/")
    assert doc.needs_ocr is False
    assert "عربي" in doc.pages[0].text


def test_load_markdown_file():
    text = "# عنوان\n\nمحتوى المستند للاختبار."
    payload = base64.b64encode(text.encode("utf-8")).decode("ascii")
    doc = load_document(payload, filename="note.md")
    assert doc.needs_ocr is False
    assert "عنوان" in doc.pages[0].text


def test_load_docx_file():
    from docx import Document

    document = Document()
    document.add_heading("سيرة ذاتية", level=1)
    document.add_paragraph("مهندس برمجيات متخصص في الذكاء الاصطناعي.")
    buf = BytesIO()
    document.save(buf)
    payload = base64.b64encode(buf.getvalue()).decode("ascii")
    doc = load_document(payload, filename="cv.docx")
    assert "wordprocessingml" in doc.mime_type
    assert doc.needs_ocr is False
    assert "ذكاء" in doc.pages[0].text or "سيرة" in doc.pages[0].text


def test_unsupported_extension_raises():
    payload = base64.b64encode(b"not-a-real-office-file").decode("ascii")
    with pytest.raises(LoaderError, match="Unsupported file type"):
        load_document(payload, filename="data.xlsx")
